"""services/resume_document.py + agents/skills/doc_generation.py 测试

覆盖：
- 模板生成路径（无上传文件）
- 原 docx 对齐改写路径（保留图标段落、替换对齐文本、插入新增行）
- 空文本、文件名消毒、同名序号
- ResumeDocxSkill 排版（节标题加粗、bullet）
"""
import os

import pytest
from docx import Document

from services.resume_document import (
    generate_resume_docx,
    sanitize_filename,
    _sanitize_base,
)
from agents.skills.doc_generation import ResumeDocxSkill


SAMPLE_TEXT = """张三
工作经历
- 2020年至今 某科技公司 高级工程师
2018-2020 某互联网公司 开发工程师
- 负责订单系统设计，QPS 提升 5 倍
教育经历
2014-2018 某大学 计算机科学"""


@pytest.fixture
def tmp_export(tmp_path, monkeypatch):
    """把导出目录重定向到临时目录，避免污染系统 temp"""
    import services.resume_document as mod
    monkeypatch.setattr(mod.tempfile, "gettempdir", lambda: str(tmp_path))
    return tmp_path


# ---------- 模板生成 ----------

def test_template_generation_basic(tmp_export):
    result = generate_resume_docx(SAMPLE_TEXT)
    assert result.mode == "template"
    assert os.path.exists(result.path)
    assert result.filename == "简历_优化版.docx"

    doc = Document(result.path)
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert texts[0] == "张三"
    assert any("订单系统" in t for t in texts)
    # bullet 符号被统一为 •
    assert any(t.startswith("•") for t in texts)


def test_empty_text_raises(tmp_export):
    with pytest.raises(ValueError):
        generate_resume_docx("   ")


def test_output_name_uses_sanitized_source_name(tmp_export):
    result = generate_resume_docx(SAMPLE_TEXT, source_name="我的简历.docx")
    # 无文件落盘 → 仍走模板；名称取消毒后的主体
    assert result.filename == "我的简历_优化版.docx"


def test_same_name_gets_sequence_suffix(tmp_export):
    first = generate_resume_docx(SAMPLE_TEXT)
    second = generate_resume_docx(SAMPLE_TEXT)
    assert first.path != second.path
    assert second.filename == "简历_优化版_1.docx"


# ---------- 文件名消毒 ----------

def test_sanitize_filename_strips_path():
    assert sanitize_filename("../../evil.docx") == "evil.docx"
    # basename 去掉目录成分，结果不含分隔符
    name = sanitize_filename("a/b\\c.docx")
    assert "/" not in name and "\\" not in name
    assert name == "c.docx"


def test_sanitize_filename_windows_special_chars():
    name = sanitize_filename('re:sume?.docx')
    for ch in '\\/:*?"<>|':
        assert ch not in name


def test_sanitize_filename_empty_default():
    assert sanitize_filename(None) == "file"
    assert sanitize_filename("") == "file"


def test_sanitize_base():
    assert _sanitize_base("../../a b.docx") == "a b"
    assert _sanitize_base(None) is None


# ---------- 对齐改写 ----------

def _build_source_docx(path):
    """构造带图标段落与粗体样式的原始简历"""
    doc = Document()
    doc.add_paragraph("李四")
    p_icon = doc.add_paragraph()
    p_icon.add_run("📌 联系方式保持不动行")
    p_work = doc.add_paragraph()
    run = p_work.add_run("2020年至今 某科技公司 负责订单系统")
    run.bold = True
    doc.add_paragraph("教育经历：某大学")
    doc.save(path)


def test_rewrite_preserves_icon_and_replaces_aligned(tmp_export):
    src = os.path.join(str(tmp_export), "source.docx")
    _build_source_docx(src)

    polished = (
        "李四\n"
        "2020年至今 某科技公司 主导订单中台架构升级，QPS提升五倍\n"
        "教育经历：某大学 计算机专业\n"
        "新增技能：Kubernetes、云原生架构"
    )
    result = generate_resume_docx(polished, source_path=src, source_name="source.docx")
    assert result.mode == "rewrite"
    assert result.filename == "source_优化版.docx"

    doc = Document(result.path)
    texts = [p.text for p in doc.paragraphs]

    # 图标行未对齐 → 原文保留（不丢图标）
    assert any("📌 联系方式保持不动行" == t for t in texts)
    # 对齐段落被改写
    assert any("主导订单中台架构升级" in t for t in texts)
    assert not any(t == "2020年至今 某科技公司 负责订单系统" for t in texts)
    # 未对齐的润色新行被追加
    assert any("Kubernetes" in t for t in texts)


def test_rewrite_keeps_paragraph_bold(tmp_export):
    src = os.path.join(str(tmp_export), "source_bold.docx")
    _build_source_docx(src)
    polished = "李四\n2020年至今 某科技公司 主导订单中台架构升级，QPS提升五倍\n教育经历：某大学"
    result = generate_resume_docx(polished, source_path=src, source_name="source_bold.docx")

    doc = Document(result.path)
    target = next(
        p for p in doc.paragraphs if "主导订单中台架构升级" in p.text
    )
    # 替换沿用首个 run 的粗体格式
    assert target.runs and target.runs[0].bold is True


def test_rewrite_corrupt_source_falls_back_to_template(tmp_export):
    bad = os.path.join(str(tmp_export), "broken.docx")
    with open(bad, "wb") as f:
        f.write(b"not a real docx zip content")
    result = generate_resume_docx(SAMPLE_TEXT, source_path=bad, source_name="broken.docx")
    assert result.mode == "template"


# ---------- ResumeDocxSkill ----------

def test_skill_section_heading_bold_and_border(tmp_path):
    out = str(tmp_path / "skill.docx")
    ResumeDocxSkill().generate(SAMPLE_TEXT, out)
    doc = Document(out)

    heading = next(p for p in doc.paragraphs if p.text == "工作经历")
    assert heading.runs[0].bold is True
    # 节标题底部边框
    pPr = heading._p.pPr
    assert pPr is not None and pPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr"
    ) is not None


def test_skill_raises_on_empty(tmp_path):
    out = str(tmp_path / "empty.docx")
    with pytest.raises(ValueError):
        ResumeDocxSkill().generate("  \n  ", out)
